"""
Fetch your CMISGo weekly timetable using the saved login session
(created by login_setup.py) and print a clean weekly report.

No visible browser window needed for normal runs (headless). If the
saved session has expired, it prints a message telling you to rerun
login_setup.py.

Usage:
    python get_timetable.py                 # current week
    python get_timetable.py --offset 1      # next week
    python get_timetable.py --offset -1 --csv   # last week, also write CSV
    python get_timetable.py --month 2026-10 --csv  # a whole calendar month
"""

import argparse
import csv
import json
import re
import sys
from collections import OrderedDict
from datetime import datetime, timedelta
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

HERE = Path(__file__).parent
STATE_FILE = HERE / "storage_state.json"
TIMETABLE_URL = "https://cmisgostudents.shef.ac.uk/CMISGo/Web/Timetable"
APPOINTMENTS_RE = re.compile(r"/Service/Appointments/Month/(\d+)/[^?]*(\?.*)?$", re.I)

# The site's own "Month/<id>/.." endpoint windows aren't linear per-week, so
# rather than compute an id for a target date, we scan a neighbourhood of
# ids around whatever id the site itself used on page load and merge
# everything. Small/fast requests, so scanning wide is cheap and robust.
SCAN_BACK = 2
SCAN_FORWARD = 8
SCAN_TIMEOUT_MS = 5000


def parse_dt(s: str) -> datetime:
    base = s.split(".")[0]
    return datetime.strptime(base, "%Y-%m-%dT%H:%M:%S")


def _merge_events(events, data):
    for ev in data:
        eid = ev.get("id")
        if eid not in events:
            ev = dict(ev)
            ev["instances"] = set(ev.get("instances", []))
            events[eid] = ev
        else:
            events[eid]["instances"].update(ev.get("instances", []))


def fetch_events(wide=False):
    if not STATE_FILE.exists():
        print("No saved session found. Run: python login_setup.py")
        sys.exit(1)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(storage_state=str(STATE_FILE))
        page = context.new_page()

        captured = []  # list of (anchor_id, query_string, body)

        def on_response(response):
            m = APPOINTMENTS_RE.search(response.url)
            if m:
                try:
                    captured.append((int(m.group(1)), m.group(2) or "", response.text()))
                except Exception:
                    pass

        page.on("response", on_response)
        page.goto(TIMETABLE_URL, wait_until="domcontentloaded")
        page.wait_for_timeout(8000)

        final_url = page.url

        if "login.shef.ac.uk" in final_url or "cas/login" in final_url:
            context.close()
            browser.close()
            print(f"Session expired or not logged in (landed on: {final_url}). Run: python login_setup.py")
            sys.exit(1)

        if not captured:
            context.close()
            browser.close()
            print("No timetable data captured. The page structure may have changed,")
            print("or the session is stale. Try: python login_setup.py")
            sys.exit(1)

        events = {}
        for _, _, body in captured:
            try:
                _merge_events(events, json.loads(body))
            except json.JSONDecodeError:
                continue

        if wide:
            anchor_id, query_string, _ = captured[0]
            for offset in range(-SCAN_BACK, SCAN_FORWARD + 1):
                mid = anchor_id + offset
                if mid == anchor_id:
                    continue
                try:
                    resp = page.request.get(
                        f"https://cmisgostudents.shef.ac.uk/CMISGo/Web/Service/Appointments/Month/{mid}/0/1{query_string}",
                        headers={"Accept": "application/json"},
                        timeout=SCAN_TIMEOUT_MS,
                    )
                    if resp.status == 200:
                        _merge_events(events, resp.json())
                except Exception:
                    continue

        context.close()
        browser.close()

        for ev in events.values():
            ev["instances"] = sorted(ev["instances"])
        return list(events.values())


def build_range_rows(events, range_start, range_end):
    rows = []
    for ev in events:
        modules = ev.get("modules") or []
        module_name = "; ".join(dict.fromkeys(m.get("name", "") for m in modules)) or ev.get("subject", "")
        module_types = ev.get("moduleTypes") or []
        event_kind = "; ".join(dict.fromkeys(mt.get("name", "") for mt in module_types)) or ev.get("eventType", "")
        duration = ev.get("duration", 0)
        location = ev.get("location", "")
        lecturer = ev.get("lecturer", "")

        for inst in ev.get("instances", []):
            start = parse_dt(inst)
            if range_start <= start < range_end:
                end = start + timedelta(minutes=duration)
                rows.append({
                    "date": start.strftime("%Y-%m-%d"),
                    "day": start.strftime("%A"),
                    "start": start.strftime("%H:%M"),
                    "end": end.strftime("%H:%M"),
                    "module": module_name,
                    "type": event_kind,
                    "location": location,
                    "lecturer": lecturer,
                    "_sort": start,
                })
    rows.sort(key=lambda r: r["_sort"])
    for r in rows:
        del r["_sort"]
    return rows


def print_report(rows, range_start, range_end, label="Timetable"):
    title = f"{label}: {range_start.strftime('%d %b %Y')} - {(range_end - timedelta(days=1)).strftime('%d %b %Y')}"
    print(title)
    print("=" * len(title))
    if not rows:
        print("No events found for this period.")
        return
    current_day = None
    for r in rows:
        if r["day"] != current_day:
            current_day = r["day"]
            print(f"\n{r['day']} {r['date']}")
            print("-" * 40)
        print(f"  {r['start']}-{r['end']}  {r['module']}")
        print(f"           {r['type']} | {r['location']}")
        if r["lecturer"]:
            print(f"           {r['lecturer']}")


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["date", "day", "start", "end", "module", "type", "location", "lecturer"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"\nSaved CSV: {path}")


def write_xlsx(rows, path, title):
    """Write a grouped, readable Excel timetable with merged week headings."""
    title_fill = PatternFill("solid", fgColor="1F4E78")
    week_fill = PatternFill("solid", fgColor="2F75B5")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    alternate_fill = PatternFill("solid", fgColor="F5F9FD")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    columns = [
        ("Date", lambda row: datetime.strptime(row["date"], "%Y-%m-%d").strftime("%A, %d %B %Y")),
        ("Time", lambda row: f"{row['start']} - {row['end']}"),
        ("Module", lambda row: row["module"]),
        ("Activity", lambda row: row["type"]),
        ("Location", lambda row: row["location"]),
        ("Lecturer(s)", lambda row: row["lecturer"]),
    ]

    weeks = OrderedDict()
    for row in rows:
        date = datetime.strptime(row["date"], "%Y-%m-%d")
        monday = date - timedelta(days=date.weekday())
        weeks.setdefault(monday, []).append(row)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timetable"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    cell = sheet["A1"]
    cell.value = title
    cell.fill = title_fill
    cell.font = Font(color="FFFFFF", bold=True, size=16)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    current_row = 3
    for week_number, (monday, events) in enumerate(weeks.items(), start=1):
        sunday = monday + timedelta(days=6)
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        cell = sheet.cell(current_row, 1)
        cell.value = f"Week {week_number} | {monday:%d %b} - {sunday:%d %b %Y}"
        cell.fill = week_fill
        cell.font = white_bold
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[current_row].height = 22
        current_row += 1

        for column, (heading, _) in enumerate(columns, start=1):
            cell = sheet.cell(current_row, column, heading)
            cell.fill = header_fill
            cell.font = Font(bold=True, color="17365D")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        current_row += 1
        first_event_row = current_row

        for index, event in enumerate(events):
            for column, (_, value) in enumerate(columns, start=1):
                cell = sheet.cell(current_row, column, value(event))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
                if index % 2:
                    cell.fill = alternate_fill
            sheet.row_dimensions[current_row].height = 42
            current_row += 1

        if events:
            sheet.row_dimensions.group(first_event_row, current_row - 1, outline_level=1, hidden=False)
        sheet.row_breaks.append(Break(id=current_row))
        current_row += 1

    for column, width in {"A": 25, "B": 15, "C": 48, "D": 22, "E": 50, "F": 34}.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = "1:1"
    workbook.save(path)
    print(f"\nSaved formatted Excel workbook: {path}")


def main():
    parser = argparse.ArgumentParser(description="Fetch CMISGo timetable report")
    parser.add_argument("--offset", type=int, default=0, help="Weeks from current week (0=this week, 1=next week)")
    parser.add_argument("--month", type=str, default=None, help="Calendar month as YYYY-MM, e.g. 2026-10")
    parser.add_argument("--from", dest="start_date", type=str, default=None, help="Start date as YYYY-MM-DD; use with --to")
    parser.add_argument("--to", dest="end_date", type=str, default=None, help="End date as YYYY-MM-DD (inclusive); use with --from")
    parser.add_argument("--csv", action="store_true", help="Also write a CSV file")
    parser.add_argument("--out", type=str, default=None, help="Exact output path for the CSV (overrides --csv default location)")
    parser.add_argument("--xlsx", action="store_true", help="Also write a formatted Excel timetable with merged week headings")
    parser.add_argument("--out-xlsx", type=str, default=None, help="Exact output path for the Excel workbook (implies --xlsx)")
    args = parser.parse_args()

    if bool(args.start_date) != bool(args.end_date):
        parser.error("--from and --to must be used together")

    if args.start_date:
        try:
            range_start = datetime.strptime(args.start_date, "%Y-%m-%d")
            last_day = datetime.strptime(args.end_date, "%Y-%m-%d")
        except ValueError:
            parser.error("--from and --to must use YYYY-MM-DD")
        if last_day < range_start:
            parser.error("--to must be on or after --from")
        range_end = last_day + timedelta(days=1)
        label = "Custom Range Timetable"
        default_name = f"timetable_{range_start.strftime('%Y-%m-%d')}_to_{last_day.strftime('%Y-%m-%d')}.csv"
    elif args.month:
        year, month = (int(x) for x in args.month.split("-"))
        range_start = datetime(year, month, 1)
        range_end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
        label = range_start.strftime("%B %Y") + " Timetable"
        default_name = f"timetable_{range_start.strftime('%Y-%m')}.csv"
    else:
        today = datetime.now()
        monday = today - timedelta(days=today.weekday())
        monday = monday.replace(hour=0, minute=0, second=0, microsecond=0)
        range_start = monday + timedelta(weeks=args.offset)
        range_end = range_start + timedelta(days=7)
        label = "Weekly Timetable"
        default_name = f"timetable_{range_start.strftime('%Y-%m-%d')}.csv"

    wide = bool(args.month) or bool(args.start_date) or abs(args.offset) > 1
    events = fetch_events(wide=wide)
    rows = build_range_rows(events, range_start, range_end)
    print_report(rows, range_start, range_end, label=label)

    if args.csv or args.out:
        if args.out:
            csv_path = Path(args.out)
            csv_path.parent.mkdir(parents=True, exist_ok=True)
        else:
            out_dir = HERE / "reports"
            out_dir.mkdir(exist_ok=True)
            csv_path = out_dir / default_name
        write_csv(rows, csv_path)

    if args.xlsx or args.out_xlsx:
        if args.out_xlsx:
            xlsx_path = Path(args.out_xlsx)
            xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        else:
            out_dir = HERE / "reports"
            out_dir.mkdir(exist_ok=True)
            xlsx_path = out_dir / default_name.replace(".csv", "_formatted.xlsx")
        write_xlsx(rows, xlsx_path, label)


if __name__ == "__main__":
    main()
